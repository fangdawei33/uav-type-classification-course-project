"""Editable Office Math for equations used by the executed classification code."""
from copy import deepcopy
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.enum.text import WD_TAB_ALIGNMENT


def node(tag, *children):
    e = OxmlElement("m:" + tag)
    for child in children:
        if isinstance(child, (list,tuple)):
            for item in child:
                e.append(run(item) if isinstance(item,str) else deepcopy(item))
        else:
            e.append(run(child) if isinstance(child, str) else deepcopy(child))
    return e


def run(value):
    r = OxmlElement("m:r")
    prop = OxmlElement("m:rPr")
    style = OxmlElement("m:sty"); style.set(qn("m:val"), "p"); prop.append(style)
    r.append(prop)
    t = OxmlElement("m:t"); t.text = str(value); r.append(t)
    return r


def sub(base, index): return node("sSub", node("e", base), node("sub", index))
def sup(base, exponent): return node("sSup", node("e", base), node("sup", exponent))
def frac(top, bottom): return node("f", node("num", top), node("den", bottom))
def seq(*parts): return parts


def summation(index, upper, value):
    p = OxmlElement("m:naryPr")
    c = OxmlElement("m:chr"); c.set(qn("m:val"), "∑"); p.append(c)
    loc = OxmlElement("m:limLoc"); loc.set(qn("m:val"), "undOvr"); p.append(loc)
    return node("nary", p, node("sub", index), node("sup", upper), node("e", value))


def root(value):
    p=OxmlElement("m:radPr"); hide=OxmlElement("m:degHide"); hide.set(qn("m:val"),"1");p.append(hide)
    return node("rad",p,node("deg"),node("e",value))


def equation_parts(key):
    # All objectives below mirror features.py, uav12_models.py and uav12_qwen.py.
    if key == "split":
        return [sub("D","tr"), " ∩ ",sub("D","te")," = ∅,   ",sub("G","tr")," ∩ ",sub("G","te")," = ∅"]
    if key == "hand":
        return [sub("x","H")," = [",sub("h","HOG"),"; ",sub("h","HSV"),"; ",sub("h","LBP"),"; q] ∈ ",sup("ℝ","1831")]
    if key == "siglip":
        return [sub("x","S")," = ",frac(seq(sub("f","θ"),"(I)"),seq("‖",sub("f","θ"),"(I)‖₂"))," ∈ ",sup("ℝ","768"),",   θ frozen"]
    if key == "scale":
        return [sub("z","j")," = ",frac(seq(sub("x","j")," − ",sub("μ","j,tr")),sub("σ","j,tr")),",   ",sub("μ","j,tr"),", ",sub("σ","j,tr")," fitted within each fold"]
    if key == "lora":
        return ["W′ = ",sub("W","NF4")," + ",frac("α","r"),"BA,   r = 8,   α = 16"]
    if key == "answer":
        return [sub("ℓ","i")," = −",frac("1",seq("|",sub("A","i"),"|")),summation(seq("t ∈ ",sub("A","i")),"",seq("log ",sub("p","φ"),"(",sub("a","i,t")," | ",sub("I","i"),", P, ",sub("a","i,<t"),")"))]
    if key == "weight":
        return [sub("u","c")," = ",root(frac("N",seq("K",sub("n","c")))),",   ",sub("w","c")," = ",frac(sub("u","c"),summation("d=0","K−1",seq(frac(sub("n","d"),"N"),sub("u","d"))))]
    if key == "objective":
        return ["ℒ = ",frac("1","B"),summation("i=1","B",seq(sub("w",sub("y","i")),sub("ℓ","i"))),",   ",sub("w","c")," = 1  (standard)"]
    if key == "macro":
        return ["Macro-F1(S) = ",frac("1","|S|"),summation("c ∈ S","",frac(seq("2",sub("P","c"),sub("R","c")),seq(sub("P","c")," + ",sub("R","c"))))]
    if key == "selection":
        return ["λ* = arg max",sub(" ","λ ∈ Λ")," ",frac("1","2"),summation("f=1","2",seq("Macro-F1(",sub("S","CV"),"; λ, f)"))]
    raise ValueError(key)


LATEX = {
    "split":r"\mathcal D_{tr}\cap\mathcal D_{te}=\varnothing,\quad G_{tr}\cap G_{te}=\varnothing",
    "hand":r"x_H=[h_{HOG};h_{HSV};h_{LBP};q]\in\mathbb R^{1831}",
    "siglip":r"x_S=\frac{f_\theta(I)}{\|f_\theta(I)\|_2}\in\mathbb R^{768},\quad\theta\text{ frozen}",
    "scale":r"z_j=\frac{x_j-\mu_{j,tr}}{\sigma_{j,tr}}",
    "lora":r"W'=W_{NF4}+\frac{\alpha}{r}BA,\quad r=8,\quad\alpha=16",
    "answer":r"\ell_i=-\frac1{|A_i|}\sum_{t\in A_i}\log p_\phi(a_{i,t}\mid I_i,P,a_{i,<t})",
    "weight":r"u_c=\sqrt{\frac{N}{Kn_c}},\quad w_c=\frac{u_c}{\sum_{d=0}^{K-1}(n_d/N)u_d}",
    "objective":r"\mathcal L=\frac1B\sum_{i=1}^B w_{y_i}\ell_i,\quad w_c=1\text{ (standard)}",
    "macro":r"MacroF1(S)=\frac1{|S|}\sum_{c\in S}\frac{2P_cR_c}{P_c+R_c}",
    "selection":r"\lambda^*=\arg\max_{\lambda\in\Lambda}\frac12\sum_{f=1}^2MacroF1(S_{CV};\lambda,f)"}


def add_equation(doc, key, number):
    p=doc.add_paragraph(); pf=p.paragraph_format
    pf.first_line_indent=Cm(0);pf.keep_together=True;pf.space_before=Pt(6);pf.space_after=Pt(7)
    pf.tab_stops.add_tab_stop(Cm(7.75),WD_TAB_ALIGNMENT.CENTER)
    pf.tab_stops.add_tab_stop(Cm(15.5),WD_TAB_ALIGNMENT.RIGHT)
    p.add_run("\t")
    math=node("oMath",*equation_parts(key));p._p.append(math)
    p.add_run("\t("+str(number)+")")
    return p
