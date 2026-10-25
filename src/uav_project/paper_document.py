"""Render the course report Markdown as an academic DOCX without running models."""
from pathlib import Path
import re

from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT


def _font(run, size=12, bold=False, chinese="宋体", latin="Times New Roman"):
    run.font.name = latin
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    rf = run._r.get_or_add_rPr().get_or_add_rFonts()
    for key, value in [("ascii", latin), ("hAnsi", latin), ("eastAsia", chinese)]:
        rf.set(qn("w:" + key), value)


def _text(p, value, size=12, bold=False, chinese="宋体"):
    # Markdown emphasis/code markers do not become literal punctuation in Word.
    for part in re.split(r"(\*\*.*?\*\*|`[^`]+`)", value):
        strong = part.startswith("**") and part.endswith("**")
        part = part[2:-2] if strong else part.strip("`")
        _font(p.add_run(part), size=size, bold=bold or strong, chinese=chinese)


def _link(p, url):
    rid = p.part.relate_to(url, RT.HYPERLINK, is_external=True)
    node = OxmlElement("w:hyperlink")
    node.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    prop = OxmlElement("w:rPr")
    for name, val in [("color", "0563C1"), ("u", "single"), ("sz", "21")]:
        el = OxmlElement("w:" + name)
        el.set(qn("w:val"), val)
        prop.append(el)
    run.append(prop)
    t = OxmlElement("w:t"); t.text = url
    run.append(t); node.append(run); p._p.append(node)


def _table(doc, lines):
    rows = [[x.strip() for x in line.strip().strip("|").split("|")] for line in lines]
    rows.pop(1)  # Markdown alignment row.
    n = len(rows[0])
    table = doc.add_table(rows=0, cols=n)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    # Give descriptive columns more room than short numeric values.
    if n == 6:
        weights = [3.5, 1.7, 1.8, 1.8, 1.8, 2.2]
    elif n == 5:
        weights = [3.4, 2.6, 2.6, 2.0, 2.0]
    elif n == 4:
        weights = [3.6, 4.3, 3.0, 3.2]
    elif n == 3:
        weights = [3.4, 4.6, 6.0]
    else:
        weights = [4.4, 10.0] if n == 2 else [1] * n
    widths = [Cm(16 * v / sum(weights)) for v in weights]
    for c, w in zip(table.columns, widths): c.width = w
    for ri, values in enumerate(rows):
        row = table.add_row()
        trpr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit"); trpr.append(cant_split)
        if ri == 0:
            trpr.append(OxmlElement("w:tblHeader"))
        for ci, (cell, value) in enumerate(zip(row.cells, values)):
            cell.width = widths[ci]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            mar = OxmlElement("w:tcMar")
            for edge, amount in [("top", 70), ("bottom", 70), ("start", 100), ("end", 100)]:
                node = OxmlElement("w:" + edge); node.set(qn("w:w"), str(amount)); node.set(qn("w:type"), "dxa"); mar.append(node)
            tcpr.append(mar)
            borders = OxmlElement("w:tcBorders")
            for edge in ["top", "bottom", "left", "right"]:
                node = OxmlElement("w:" + edge)
                active = (ri == 0 and edge in ["top", "bottom"]) or (ri == len(rows)-1 and edge == "bottom")
                node.set(qn("w:val"), "single" if active else "nil")
                if active:
                    node.set(qn("w:sz"), "12" if edge == "top" or ri == len(rows)-1 else "6")
                    node.set(qn("w:color"), "000000")
                borders.append(node)
            tcpr.append(borders)
            p = cell.paragraphs[0]
            pf = p.paragraph_format
            pf.first_line_indent = Cm(0); pf.space_after = Pt(0); pf.space_before = Pt(0)
            pf.line_spacing = 1.12
            # Entire compact tables move as a unit; long tables repeat their headers.
            pf.keep_with_next = len(rows) <= 16 and ri < len(rows)-1
            numeric = bool(re.fullmatch(r"[0-9+−.,%± /—–-]+", value))
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if ri == 0 or numeric else WD_ALIGN_PARAGRAPH.LEFT
            _text(p, value, size=9.5 if n >= 5 else 10, bold=(ri == 0))
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 0.35; p.paragraph_format.first_line_indent = Cm(0)


def render_paper(markdown, destination):
    markdown, destination = Path(markdown), Path(destination)
    doc = Document()
    # Bundled Word defaults may carry a blue rule in the Title style.
    for borders in list(doc.styles.element.iter(qn("w:pBdr"))):
        borders.getparent().remove(borders)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(2.54)
    sec.left_margin, sec.right_margin = Cm(2.6), Cm(2.4)
    sec.footer_distance = Cm(1.35)
    for name, size, bold, chinese in [("Normal",12,False,"宋体"),("Title",18,True,"黑体"),
                                       ("Heading 1",16,True,"黑体"),("Heading 2",14,True,"黑体"),
                                       ("Heading 3",12,True,"黑体")]:
        st = doc.styles[name]
        st.font.name = "Times New Roman"; st.font.size = Pt(size); st.font.bold = bold
        st.font.color.rgb = RGBColor(0,0,0)
        st._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), chinese)
        pf = st.paragraph_format
        pf.first_line_indent = Cm(.74 if name == "Normal" else 0)
        pf.line_spacing = 1.4
        pf.space_after = Pt(5 if name == "Normal" else 7)
        pf.space_before = Pt(12 if name.startswith("Heading") else 0)
        pf.keep_with_next = name != "Normal"
        pf.widow_control = True
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.first_line_indent = Cm(0)
    _font(footer.add_run("无人机机型图像分类与多模型比较研究   "),9)
    field = OxmlElement("w:fldSimple"); field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    lines = markdown.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line: i += 1; continue
        if line.startswith("# "):
            p = doc.add_paragraph(style="Title"); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _text(p,line[2:],18,True,"黑体")
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Cm(0)
            _text(p,"机器学习课程综合设计实验报告",14,True,"黑体")
            i += 1; continue
        if line.startswith("### ") or line.startswith("## "):
            level = 2 if line.startswith("### ") else 1
            value = line[4:] if level == 2 else line[3:]
            p = doc.add_paragraph(style=f"Heading {level}")
            _text(p,value,14 if level == 2 else 16,True,"黑体")
            if value.startswith("1 实验概述"):
                p.paragraph_format.page_break_before = True
            i += 1; continue
        if line.startswith("!["):
            path = line.split("](",1)[1].removesuffix(")")
            image = markdown.parent / path
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_after = Pt(0)
            if image.is_file():
                pic = p.add_run().add_picture(str(image),width=Cm(15.5))
                if pic.height > Cm(14):
                    factor = Cm(14)/pic.height
                    pic.width = int(pic.width*factor); pic.height = Cm(14)
            else: raise FileNotFoundError(image)
            i += 1; continue
        if line.startswith("| "):
            group=[]
            while i < len(lines) and lines[i].strip().startswith("|"):
                group.append(lines[i]); i += 1
            _table(doc, group); continue
        if line.startswith("```"):
            code=[]; i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                code.append(lines[i]); i += 1
            p = doc.add_paragraph()
            pf=p.paragraph_format; pf.first_line_indent=Cm(0); pf.left_indent=Cm(.25); pf.right_indent=Cm(.1)
            pf.line_spacing=1.03; pf.space_before=Pt(3); pf.space_after=Pt(6)
            pf.keep_together=True
            for j,value in enumerate(code):
                if j: p.add_run().add_break()
                r=p.add_run(value); _font(r,8.2,latin="Consolas")
            i += 1; continue
        p=doc.add_paragraph()
        if re.match(r"^(表\d|图\d|代码\d)",line):
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent=Cm(0)
            p.paragraph_format.keep_with_next=line.startswith(("表","代码"))
            _text(p,line,10.5,True)
        elif line.startswith("项目代码仓库："):
            p.paragraph_format.first_line_indent=Cm(0)
            _text(p,"项目代码仓库：",10.5)
            _link(p,line.split("：",1)[1].strip())
        elif line.startswith("[" ) and re.match(r"\[\d+\]",line):
            p.paragraph_format.first_line_indent=Cm(-.65); p.paragraph_format.left_indent=Cm(.65)
            _text(p,line,10.5)
        elif line.startswith("注：") or line.startswith("来源："):
            p.paragraph_format.first_line_indent=Cm(0)
            _text(p,line,10)
        else:
            _text(p,line)
        i += 1
    doc.core_properties.title = lines[0].removeprefix("# ")
    doc.core_properties.subject = "机器学习课程综合设计实验报告"
    destination.parent.mkdir(parents=True, exist_ok=True)
    doc.save(destination)
    return destination
